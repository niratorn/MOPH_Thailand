"""Render validation results and summaries for humans and machines."""

from __future__ import annotations

import csv
import json
from collections import Counter
from typing import TextIO

from .loader import DataFile
from .schema import Schema
from .validator import ValidationResult


def print_validation_summary(result: ValidationResult, out: TextIO) -> None:
    """Write a concise human-readable summary to ``out``."""
    status = "PASS" if result.is_valid else "FAIL"
    out.write(f"Schema:        {result.schema_name}\n")
    out.write(f"Rows checked:  {result.total_rows}\n")
    out.write(f"Errors:        {len(result.errors)} "
              f"(in {result.rows_with_errors()} rows)\n")
    out.write(f"Warnings:      {len(result.warnings)}\n")
    out.write(f"Result:        {status}\n")

    if result.issues:
        by_rule = Counter(i.rule for i in result.issues)
        out.write("\nFindings by rule:\n")
        for rule, count in by_rule.most_common():
            out.write(f"  {rule:<22} {count}\n")

    # Show the first handful of issues inline so a user gets immediate feedback.
    preview = result.issues[:10]
    if preview:
        out.write("\nFirst findings:\n")
        for i in preview:
            loc = "header" if i.row == 0 else f"row {i.row}"
            out.write(f"  [{i.severity}] {loc} {i.column}: {i.message}\n")
        if len(result.issues) > len(preview):
            out.write(f"  ... and {len(result.issues) - len(preview)} more "
                      f"(use --out to write the full report)\n")


def write_issues_csv(result: ValidationResult, path: str) -> None:
    """Write every issue to a CSV report for review/correction."""
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["row", "column", "severity", "rule", "message", "value"])
        for i in result.issues:
            writer.writerow([i.row, i.column, i.severity, i.rule, i.message, i.value])


def result_to_dict(result: ValidationResult) -> dict:
    return {
        "schema": result.schema_name,
        "rows_checked": result.total_rows,
        "error_count": len(result.errors),
        "warning_count": len(result.warnings),
        "rows_with_errors": result.rows_with_errors(),
        "is_valid": result.is_valid,
        "issues": [
            {
                "row": i.row,
                "column": i.column,
                "severity": i.severity,
                "rule": i.rule,
                "message": i.message,
                "value": i.value,
            }
            for i in result.issues
        ],
    }


def write_result_json(result: ValidationResult, path: str) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(result_to_dict(result), fh, ensure_ascii=False, indent=2)


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return openpyxl
    except ImportError as exc:  # pragma: no cover - exercised via CLI message
        raise SystemExit(
            "Excel output requires the optional 'openpyxl' package.\n"
            "Install it with:  pip install openpyxl   (or:  pip install "
            "'moph-report[excel]')"
        ) from exc


def write_issues_xlsx(result: ValidationResult, path: str) -> None:
    """Write the validation result to a two-sheet Excel workbook.

    Sheet 1 ("Summary") gives the headline counts; sheet 2 ("Issues") lists
    every finding. Friendlier for non-technical staff than a raw CSV.
    """
    openpyxl = _require_openpyxl()
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Schema", result.schema_name])
    ws.append(["Rows checked", result.total_rows])
    ws.append(["Errors", len(result.errors)])
    ws.append(["Warnings", len(result.warnings)])
    ws.append(["Rows with errors", result.rows_with_errors()])
    ws.append(["Result", "PASS" if result.is_valid else "FAIL"])
    ws.append([])
    ws.append(["Findings by rule", "Count"])
    for rule, count in Counter(i.rule for i in result.issues).most_common():
        ws.append([rule, count])

    issues_ws = wb.create_sheet("Issues")
    issues_ws.append(["row", "column", "severity", "rule", "message", "value"])
    for i in result.issues:
        issues_ws.append([i.row, i.column, i.severity, i.rule, i.message, i.value])

    wb.save(path)


def write_summary_xlsx(summary: dict, path: str) -> None:
    """Write aggregate statistics to an Excel workbook (one sheet per field)."""
    openpyxl = _require_openpyxl()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Overview"
    ws.append(["Schema", summary["schema"]])
    ws.append(["Thai name", summary["thai_name"]])
    ws.append(["Total rows", summary["total_rows"]])

    for field_name, counts in summary["breakdowns"].items():
        sheet = wb.create_sheet(field_name[:31])  # Excel sheet-name limit
        sheet.append([field_name, "Count"])
        for code, n in counts.items():
            sheet.append([code, n])

    wb.save(path)


def build_summary(data: DataFile, schema: Schema) -> dict:
    """Produce simple aggregate statistics for a data file.

    Counts total rows and, for each coded field that defines an ``allowed`` set,
    a breakdown of how many rows carry each code. Useful for a quick monthly
    sanity check (e.g. sex distribution, diagnosis-type counts).
    """
    summary: dict = {
        "schema": schema.name,
        "thai_name": schema.thai_name,
        "total_rows": len(data.rows),
        "breakdowns": {},
    }
    coded = [f for f in schema.fields if f.allowed and f.name in data.header]
    records = list(data.as_dicts())
    for fld in coded:
        counter: Counter[str] = Counter(
            (r.get(fld.name, "") or "(blank)") for r in records
        )
        summary["breakdowns"][fld.name] = dict(counter.most_common())
    return summary


def print_summary(summary: dict, out: TextIO) -> None:
    out.write(f"Schema:     {summary['schema']} ({summary['thai_name']})\n")
    out.write(f"Total rows: {summary['total_rows']}\n")
    for field_name, counts in summary["breakdowns"].items():
        out.write(f"\n{field_name}:\n")
        for code, n in counts.items():
            out.write(f"  {code:<12} {n}\n")
